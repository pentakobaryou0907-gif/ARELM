import pandas as pd
import os
import sys
import time
from openpyxl import load_workbook
from datetime import datetime

# 統合Excel管理システムをインポート
try:
    from excel_manager import (
        init_master_excel, get_barcode_dict, add_to_barcode_master,
        get_products_sheet, add_product_to_sheet, update_summary_sheet,
        get_or_create_box_helper, get_next_product_number_helper
    )
    EXCEL_MANAGER_AVAILABLE = True
except ImportError:
    # フォールバック用の設定
    BARCODE_MASTER_FILE = "barcode_master.xlsx"
    OUTPUT_FILE = "sorted_products.xlsx"
    EXCEL_MANAGER_AVAILABLE = False
    print("⚠️  excel_manager.pyが見つかりません。統合Excelファイル機能を使用できません。")

# バーコード読み取り用ライブラリ（オプション）
try:
    from pyzbar.pyzbar import decode
    import cv2
    CAMERA_AVAILABLE = True
except ImportError:
    CAMERA_AVAILABLE = False
    print("⚠️  カメラ機能を使用するには: pip install pyzbar opencv-python")

# ==========================
# 設定
# ==========================
CATEGORY_RULES = {
    "T-shirt": "Summer",
    "Jacket": "Winter",
    "Scarf": "Winter",
    "Sunglasses": "Summer",
    "Socks": "00"
}
BOX_CAPACITY = 5  # 1箱あたりの最大数


# ==========================
# 箱クラス
# ==========================
class Box:
    def __init__(self, prefix, number, suffix, capacity):
        self.prefix = prefix
        self.number = number
        self.suffix = suffix
        self.capacity = capacity
        self.items = []

    def add_item(self, item):
        if len(self.items) < self.capacity:
            self.items.append(item)
            return True
        return False

    def is_full(self):
        return len(self.items) >= self.capacity

    def label(self):
        return f"{self.prefix}-{self.number}{self.suffix}"


# ==========================
# バーコードマスター管理
# ==========================
def load_barcode_master():
    """バーコードマスターファイルを読み込む"""
    if EXCEL_MANAGER_AVAILABLE:
        # 統合Excelファイルを使用
        init_master_excel()
        return get_barcode_dict()
    else:
        # フォールバック: 古い方法
        if not os.path.exists(BARCODE_MASTER_FILE):
            df = pd.DataFrame(columns=["バーコード", "商品名"])
            df.to_excel(BARCODE_MASTER_FILE, index=False)
            print(f"📝 新規作成: {BARCODE_MASTER_FILE}")
            return {}
        
        try:
            df = pd.read_excel(BARCODE_MASTER_FILE)
            if "バーコード" not in df.columns or "商品名" not in df.columns:
                print(f"❌ エラー: {BARCODE_MASTER_FILE} の列名が正しくありません。")
                return {}
            
            barcode_dict = {}
            for _, row in df.iterrows():
                barcode = str(row["バーコード"]).strip()
                product_name = str(row["商品名"]).strip()
                if barcode and product_name:
                    barcode_dict[barcode] = product_name
            
            print(f"✅ バーコードマスターを読み込みました: {len(barcode_dict)}件")
            return barcode_dict
        except Exception as e:
            print(f"❌ エラー: バーコードマスターの読み込みに失敗しました: {str(e)}")
            return {}


def add_to_barcode_master_local(barcode, product_name):
    """バーコードマスターに新しいエントリを追加"""
    if EXCEL_MANAGER_AVAILABLE:
        # 統合Excelファイルを使用
        from excel_manager import add_to_barcode_master as excel_add_to_barcode_master
        return excel_add_to_barcode_master(barcode, product_name)
    else:
        # フォールバック: 古い方法
        try:
            if os.path.exists(BARCODE_MASTER_FILE):
                df = pd.read_excel(BARCODE_MASTER_FILE)
            else:
                df = pd.DataFrame(columns=["バーコード", "商品名"])
            
            if barcode in df["バーコード"].astype(str).values:
                print(f"⚠️  バーコード {barcode} は既に登録されています。")
                return False
            
            new_row = pd.DataFrame({"バーコード": [barcode], "商品名": [product_name]})
            df = pd.concat([df, new_row], ignore_index=True)
            df.to_excel(BARCODE_MASTER_FILE, index=False)
            print(f"✅ バーコードマスターに追加しました: {barcode} → {product_name}")
            return True
        except Exception as e:
            print(f"❌ エラー: バーコードマスターへの追加に失敗しました: {str(e)}")
            return False


# ==========================
# 既存のExcelファイル管理
# ==========================
def load_existing_data():
    """既存の出力ファイルを読み込む"""
    if not os.path.exists(OUTPUT_FILE):
        return pd.DataFrame(columns=["商品名", "商品番号", "箱名", "バーコード", "登録日時"])
    
    try:
        df = pd.read_excel(OUTPUT_FILE)
        # 必要な列がない場合は追加
        required_columns = ["商品名", "商品番号", "箱名", "バーコード", "登録日時"]
        for col in required_columns:
            if col not in df.columns:
                df[col] = ""
        return df
    except Exception as e:
        print(f"⚠️  既存ファイルの読み込みエラー: {str(e)}")
        return pd.DataFrame(columns=["商品名", "商品番号", "箱名", "バーコード", "登録日時"])


def get_next_product_number(existing_df):
    """次の商品番号を取得"""
    if existing_df.empty or "商品番号" not in existing_df.columns:
        return 1
    
    # 既存の商品番号から最大値を取得
    max_num = 0
    for num_str in existing_df["商品番号"].astype(str):
        try:
            num = int(num_str)
            max_num = max(max_num, num)
        except:
            pass
    
    return max_num + 1


def get_or_create_box(boxes, product_name):
    """商品を適切な箱に割り当て、新しい箱が必要なら作成"""
    prefix = CATEGORY_RULES.get(product_name, "00")
    
    if prefix not in boxes:
        boxes[prefix] = [Box(prefix, 1, "A", BOX_CAPACITY)]
    
    current_box = boxes[prefix][-1]
    
    if not current_box.add_item(product_name):
        # 箱が満杯なので新しい箱を作成
        new_suffix = chr(ord(current_box.suffix) + 1)
        new_box = Box(prefix, current_box.number, new_suffix, BOX_CAPACITY)
        new_box.add_item(product_name)
        boxes[prefix].append(new_box)
        current_box = new_box
    
    return current_box


def add_product_to_excel(barcode, product_name, barcode_master):
    """バーコードを読み取ってExcelに追加"""
    try:
        if EXCEL_MANAGER_AVAILABLE:
            # 統合Excelファイルを使用
            init_master_excel()
            
            # 既存データを読み込む
            existing_df = get_products_sheet()
            
            # 箱の状態を復元
            boxes = {}
            if not existing_df.empty and "箱名" in existing_df.columns:
                # 箱名から箱の状態を復元
                for _, row in existing_df.iterrows():
                    product = row["商品名"]
                    box_name = str(row["箱名"])
                    if product and box_name:
                        prefix = CATEGORY_RULES.get(product, "00")
                        if prefix not in boxes:
                            boxes[prefix] = []
                        # 箱名を解析（例: "Summer-1A"）
                        parts = box_name.split('-')
                        if len(parts) == 2:
                            box_num = int(parts[1][0]) if parts[1][0].isdigit() else 1
                            box_suffix = parts[1][1] if len(parts[1]) > 1 else 'A'
                            box = Box(prefix, box_num, box_suffix, BOX_CAPACITY)
                            if box.label() not in [b.label() for b in boxes[prefix]]:
                                boxes[prefix].append(box)
            
            # 次の商品番号を取得
            product_number = get_next_product_number_helper()
            
            # 箱を取得または作成
            current_box = get_or_create_box_helper(boxes, product_name, CATEGORY_RULES, BOX_CAPACITY)
            
            # 商品を追加（分類情報は空欄で追加、後でExcelで編集可能）
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            success = add_product_to_sheet(product_name, product_number, current_box.label(), barcode, "", "", "", "", timestamp)
            
            if success:
                # 集計シートを更新
                update_summary_sheet()
                print(f"✅ 追加完了: {product_name} (バーコード: {barcode}) → {current_box.label()}")
            return success
        else:
            # フォールバック: 古い方法
            existing_df = load_existing_data()
            boxes = {}
            for _, row in existing_df.iterrows():
                product = row["商品名"]
                box_name = row["箱名"]
                if product and box_name:
                    prefix = CATEGORY_RULES.get(product, "00")
                    if prefix not in boxes:
                        boxes[prefix] = [Box(prefix, 1, "A", BOX_CAPACITY)]
            
            product_number = get_next_product_number(existing_df)
            current_box = get_or_create_box(boxes, product_name)
            
            new_row = pd.DataFrame({
                "商品名": [product_name],
                "商品番号": [f"{product_number:03}"],
                "箱名": [current_box.label()],
                "バーコード": [barcode],
                "登録日時": [datetime.now().strftime("%Y-%m-%d %H:%M:%S")]
            })
            
            updated_df = pd.concat([existing_df, new_row], ignore_index=True)
            
            with pd.ExcelWriter(OUTPUT_FILE, engine='openpyxl') as writer:
                updated_df.to_excel(writer, index=False, sheet_name='Sheet1')
            
            wb = load_workbook(OUTPUT_FILE)
            ws = wb.active
            for row in range(2, len(updated_df) + 2):
                cell = ws[f'B{row}']
                if cell.value:
                    cell.value = str(cell.value).zfill(3)
                    cell.number_format = '@'
            wb.save(OUTPUT_FILE)
            wb.close()
            
            print(f"✅ 追加完了: {product_name} (バーコード: {barcode}) → {current_box.label()}")
            return True
        
    except Exception as e:
        print(f"❌ エラー: Excelへの追加に失敗しました: {str(e)}")
        return False


# ==========================
# バーコード読み取り（USBリーダー）
# ==========================
def read_barcode_from_scanner():
    """USBバーコードリーダーから読み取り（キーボード入力として）"""
    if EXCEL_MANAGER_AVAILABLE:
        init_master_excel()
    
    print("\n📦 scannerモード: バーコードをスキャンしてください...")
    print("   (Enterキーを押すと終了します)\n")
    
    barcode_master = load_barcode_master()
    
    while True:
        try:
            # バーコードリーダーからの入力を待機
            barcode = input().strip()
            
            if not barcode:
                break
            
            # バーコードマスターから商品名を取得
            if barcode in barcode_master:
                product_name = barcode_master[barcode]
                print(f"📦 商品: {product_name} (バーコード: {barcode})")
                
                # Excelに追加
                add_product_to_excel(barcode, product_name, barcode_master)
            else:
                print(f"⚠️  バーコード {barcode} が見つかりません。")
                print("   商品名を入力してください:")
                product_name = input("商品名: ").strip()
                
                if product_name:
                    # バーコードマスターに追加
                    add_to_barcode_master_local(barcode, product_name)
                    # Excelに追加
                    add_product_to_excel(barcode, product_name, barcode_master)
                else:
                    print("❌ 商品名が入力されませんでした。")
        
        except KeyboardInterrupt:
            print("\n\n👋 終了します。")
            break
        except EOFError:
            break


# ==========================
# バーコード読み取り（カメラ）
# ==========================
def read_barcode_from_camera():
    """カメラを使ってバーコードを読み取る"""
    if not CAMERA_AVAILABLE:
        print("❌ カメラ機能を使用するには: pip install pyzbar opencv-python")
        return
    
    if EXCEL_MANAGER_AVAILABLE:
        init_master_excel()
    
    print("\n📷 カメラモード: バーコードをカメラに向けてください...")
    print("   (qキーを押すと終了します)\n")
    
    barcode_master = load_barcode_master()
    cap = cv2.VideoCapture(0)
    
    last_barcode = None
    last_barcode_time = 0
    
    try:
        while True:
            ret, frame = cap.read()
            if not ret:
                break
            
            # バーコードを検出
            barcodes = decode(frame)
            
            for barcode in barcodes:
                barcode_data = barcode.data.decode('utf-8')
                
                # 同じバーコードを連続で読み取らないようにする（1秒間隔）
                current_time = time.time()
                if barcode_data == last_barcode and current_time - last_barcode_time < 1.0:
                    continue
                
                last_barcode = barcode_data
                last_barcode_time = current_time
                
                # バーコードを画面に表示
                (x, y, w, h) = barcode.rect
                cv2.rectangle(frame, (x, y), (x + w, y + h), (0, 255, 0), 2)
                cv2.putText(frame, barcode_data, (x, y - 10), 
                           cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 255, 0), 2)
                
                print(f"\n📦 バーコード読み取り: {barcode_data}")
                
                # バーコードマスターから商品名を取得
                if barcode_data in barcode_master:
                    product_name = barcode_master[barcode_data]
                    print(f"   商品: {product_name}")
                    
                    # Excelに追加
                    add_product_to_excel(barcode_data, product_name, barcode_master)
                else:
                    print(f"⚠️  バーコード {barcode_data} が見つかりません。")
                    print("   商品名を入力してください (Enterでスキップ):")
                    product_name = input("商品名: ").strip()
                    
                    if product_name:
                        # バーコードマスターに追加
                        add_to_barcode_master_local(barcode_data, product_name)
                        # Excelに追加
                        add_product_to_excel(barcode_data, product_name, barcode_master)
            
            # 画面に表示
            cv2.imshow('Barcode Scanner - Press Q to quit', frame)
            
            if cv2.waitKey(1) & 0xFF == ord('q'):
                break
    
    except Exception as e:
        print(f"❌ エラー: {str(e)}")
    finally:
        cap.release()
        cv2.destroyAllWindows()


# ==========================
# メイン関数
# ==========================
def main():
    if EXCEL_MANAGER_AVAILABLE:
        init_master_excel()
        print("=" * 50)
        print("📦 バーコード読み取りシステム")
        print(f"📁 統合ファイル: 商品管理マスター.xlsx")
        print("=" * 50)
    else:
        print("=" * 50)
        print("📦 バーコード読み取りシステム")
        print("=" * 50)
    
    print("\nモードを選択してください:")
    print("1. USBバーコードリーダー (キーボード入力)")
    print("2. カメラでスキャン")
    print("3. 終了")
    
    while True:
        try:
            choice = input("\n選択 (1-3): ").strip()
            
            if choice == "1":
                read_barcode_from_scanner()
                break
            elif choice == "2":
                read_barcode_from_camera()
                break
            elif choice == "3":
                print("👋 終了します。")
                sys.exit(0)
            else:
                print("❌ 無効な選択です。1-3を入力してください。")
        
        except KeyboardInterrupt:
            print("\n\n👋 終了します。")
            sys.exit(0)


if __name__ == "__main__":
    main()

