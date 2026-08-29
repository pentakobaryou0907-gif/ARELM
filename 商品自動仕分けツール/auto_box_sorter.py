import pandas as pd
import string
import os
import sys

from openpyxl import load_workbook

# ==========================
# 1️⃣ ルール設定
# ==========================
CATEGORY_RULES = {
    "T-shirt": "Summer",
    "Jacket": "Winter",
    "Scarf": "Winter",
    "Sunglasses": "Summer",
    "Socks": "00"
}

BOX_CAPACITY = 5  # 1箱あたりの最大数


# ==========================
# 2️⃣ 箱クラス
# ==========================
class Box:
    def __init__(self, prefix, number, suffix, capacity):
        self.prefix = prefix
        self.number = number
        self.suffix = suffix
        self.capacity = capacity
        self.items = []

    def add_item(self, item):
        if len(self.items) < self.capacity:
            self.items.append(item)
            return True
        return False

    def is_full(self):
        return len(self.items) >= self.capacity

    def label(self):
        return f"{self.prefix}-{self.number}{self.suffix}"


# ==========================
# 3️⃣ 自動仕分け処理
# ==========================
def auto_sort(products):
    boxes = {}
    all_data = []
    product_counter = 1

    for product in products:
        prefix = CATEGORY_RULES.get(product, "00")  # 未定義品は「00」へ

        # まだ箱がないカテゴリなら新規作成
        if prefix not in boxes:
            boxes[prefix] = [Box(prefix, 1, "A", BOX_CAPACITY)]

        current_box = boxes[prefix][-1]

        # 満杯なら新しい箱を作る
        if not current_box.add_item(product):
            new_suffix = chr(ord(current_box.suffix) + 1)
            new_box = Box(prefix, current_box.number, new_suffix, BOX_CAPACITY)
            new_box.add_item(product)
            boxes[prefix].append(new_box)
            current_box = new_box

        # 結果をリストに追加
        all_data.append({
            "商品名": product,
            "商品番号": f"{product_counter:03}",
            "箱名": current_box.label()
        })
        product_counter += 1

    df = pd.DataFrame(all_data)
    # 商品番号を文字列として保存（先頭の0を保持するため）
    df["商品番号"] = df["商品番号"].astype(str)
    return df


# ==========================
# 4️⃣ Excel入出力処理
# ==========================
def main():
    # 統合Excelファイルを使用するか確認
    try:
        from excel_manager import (
            init_master_excel, get_products_sheet, save_products_sheet,
            update_summary_sheet, SHEET_RAW_INPUT, MASTER_EXCEL_FILE
        )
        USE_INTEGRATED_EXCEL = True
    except ImportError:
        USE_INTEGRATED_EXCEL = False
        input_file = "products.xlsx"
        output_file = "sorted_products.xlsx"
    
    if USE_INTEGRATED_EXCEL:
        # 統合Excelファイルを使用
        init_master_excel()
        print(f"📁 統合ファイルを使用: {MASTER_EXCEL_FILE}")
        
        # 入力データシートから商品を読み込む
        try:
            df = pd.read_excel(MASTER_EXCEL_FILE, sheet_name=SHEET_RAW_INPUT)
            if df.empty:
                print(f"⚠️  入力データシートが空です。")
                print(f"   {MASTER_EXCEL_FILE} の「{SHEET_RAW_INPUT}」シートに商品名を入力してください。")
                sys.exit(1)
        except Exception as e:
            print(f"❌ エラー: 入力データの読み込みに失敗しました: {str(e)}")
            sys.exit(1)
    else:
        # フォールバック: 古い方法
        input_file = "products.xlsx"
        output_file = "sorted_products.xlsx"
        
        if not os.path.exists(input_file):
            print(f"❌ エラー: {input_file} が見つかりません。")
            sys.exit(1)
        
        df = pd.read_excel(input_file)
    
    # 商品名列の存在確認
    if "商品名" not in df.columns:
        print(f"❌ エラー: Excelファイルに「商品名」列が見つかりません。")
        print(f"   見つかった列: {', '.join(df.columns)}")
        sys.exit(1)
    
    products = df["商品名"].tolist()
    
    # 空のリストチェック
    if not products:
        print(f"❌ エラー: 商品リストが空です。")
        sys.exit(1)

    try:
        # 自動仕分け実行
        sorted_df = auto_sort(products)
        
        if USE_INTEGRATED_EXCEL:
            # 統合Excelファイルの商品一覧シートに追加
            existing_df = get_products_sheet()
            
            # 既存データと統合（重複を避ける）
            if existing_df.empty:
                save_products_sheet(sorted_df)
            else:
                # 既存データと結合
                combined_df = pd.concat([existing_df, sorted_df]).drop_duplicates(
                    subset=["商品名", "商品番号", "箱名"], keep='first'
                )
                save_products_sheet(combined_df)
            
            # 集計シートを更新
            update_summary_sheet()
            
            print(f"✅ 仕分け完了！→ {MASTER_EXCEL_FILE} の「商品一覧」シートに追加しました。")
        else:
            # フォールバック: 古い方法
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                sorted_df.to_excel(writer, index=False, sheet_name='Sheet1')
            
            wb = load_workbook(output_file)
            ws = wb.active
            for row in range(2, len(sorted_df) + 2):
                cell = ws[f'B{row}']
                cell.value = str(cell.value).zfill(3)
                cell.number_format = '@'
            wb.save(output_file)
            wb.close()
            
            print(f"✅ 仕分け完了！→ {output_file} に出力しました。")
        
        print(f"   処理した商品数: {len(products)}件")
        
    except Exception as e:
        print(f"❌ エラーが発生しました: {str(e)}")
        sys.exit(1)


if __name__ == "__main__":
    main()

