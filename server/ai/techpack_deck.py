# -*- coding: utf-8 -*-
"""
テックパックスライド（.pptx）を組み立てる

なぜこれが要るのか:
    Alibaba OEMでの商品開発では、1着ごとに
    「デザイン画像＋寸法・素材・原価」をまとめた資料（テックパック）が要る。

    Canva公式APIには、この形（1着1スライド・画像とテキストの精密な
    配置）を無料で自動組み立てできる仕組みが無い
    （本命のAutofill APIはCanva Enterprise専用で、個人利用では使えない）。

    そこで、この端末だけで .pptx ファイルとして完成させる。
    .pptx は Canva・PowerPoint・Keynote のどれでも開けるので、
    Canvaで編集したいときは、Canva側の「インポート」機能（無料）で
    このファイルを取り込めばよい。外部APIへは一切送らない。
"""

import base64
import io
import os
import re
import time

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

# 16:9のスライドサイズ
スライド幅 = Cm(33.867)
スライド高さ = Cm(19.05)


def _画像データを開く(dataUrl):
    """data URL(base64) でもファイルパスでも読めるようにする。読めなければ None。"""
    if not dataUrl:
        return None
    try:
        if dataUrl.startswith('data:'):
            m = re.match(r'data:[^;]+;base64,(.+)', dataUrl, re.S)
            if not m:
                return None
            return io.BytesIO(base64.b64decode(m.group(1)))
        if os.path.isfile(dataUrl):
            return open(dataUrl, 'rb')
    except (OSError, ValueError):
        return None
    return None


def _表紙を作る(prs, タイトル):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box = slide.shapes.add_textbox(Cm(2), Cm(7), スライド幅 - Cm(4), Cm(4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = タイトル
    tf.paragraphs[0].font.size = Pt(40)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    sub = slide.shapes.add_textbox(Cm(2), Cm(11.5), スライド幅 - Cm(4), Cm(2))
    sub.text_frame.text = f'作成日: {time.strftime("%Y-%m-%d")}　Alibaba OEM 商品開発資料'
    sub.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    sub.text_frame.paragraphs[0].font.size = Pt(16)
    sub.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _商品スライドを作る(prs, 商品):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    name = (商品.get('name') or '(商品名未設定)').strip()
    description = (商品.get('description') or '').strip()
    techpack = 商品.get('techpack') or {}

    # --- 左: デザイン画像 ---
    画像 = _画像データを開く(商品.get('画像dataUrl') or 商品.get('画像path'))
    if 画像:
        try:
            slide.shapes.add_picture(画像, Cm(1), Cm(1.5), height=Cm(14))
        except Exception:
            # 壊れた画像等。無理に載せず、枠だけ残す。
            画像 = None
        finally:
            try:
                画像.close()
            except Exception:
                pass
    if not 画像:
        枠 = slide.shapes.add_textbox(Cm(1), Cm(7), Cm(13), Cm(2))
        枠.text_frame.text = '（デザイン画像なし）'
        枠.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        枠.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # --- 右上: 商品名・説明 ---
    見出し = slide.shapes.add_textbox(Cm(15.5), Cm(1), スライド幅 - Cm(17), Cm(2))
    見出し.text_frame.word_wrap = True
    見出し.text_frame.text = name
    見出し.text_frame.paragraphs[0].font.size = Pt(28)
    見出し.text_frame.paragraphs[0].font.bold = True

    if description:
        説明 = slide.shapes.add_textbox(Cm(15.5), Cm(2.8), スライド幅 - Cm(17), Cm(2.5))
        説明.text_frame.word_wrap = True
        説明.text_frame.text = description
        説明.text_frame.paragraphs[0].font.size = Pt(14)

    # --- 右中: 素材・原価・納期 ---
    情報行 = []
    if techpack.get('素材'):
        情報行.append(f"素材: {techpack['素材']}")
    if techpack.get('原価'):
        情報行.append(f"原価目安: {techpack['原価']}")
    if techpack.get('納期'):
        情報行.append(f"納期目安: {techpack['納期']}")
    if 情報行:
        情報箱 = slide.shapes.add_textbox(Cm(15.5), Cm(5.5), スライド幅 - Cm(17), Cm(3))
        tf = 情報箱.text_frame
        tf.word_wrap = True
        for i, 行 in enumerate(情報行):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = 行
            p.font.size = Pt(14)

    # --- 右下: 寸法表 ---
    寸法行たち = techpack.get('寸法行たち') or []
    if 寸法行たち:
        rows = len(寸法行たち) + 1
        table_shape = slide.shapes.add_table(
            rows, 3, Cm(15.5), Cm(9), スライド幅 - Cm(17), Cm(min(0.9 * rows, 9)))
        table = table_shape.table
        table.cell(0, 0).text = 'サイズ'
        table.cell(0, 1).text = '部位'
        table.cell(0, 2).text = '寸法(cm)'
        for r, 行 in enumerate(寸法行たち, start=1):
            table.cell(r, 0).text = str(行.get('size', ''))
            table.cell(r, 1).text = str(行.get('part', ''))
            table.cell(r, 2).text = str(行.get('value', ''))
        for row in table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(11)


def スライドを作る(商品リスト, 出力パス, タイトル='OEM商品開発資料'):
    """
    商品ごとに1スライド（デザイン画像＋テックパック情報）を作り、.pptx として保存する。

    @param 商品リスト [{name, description, 画像dataUrl or 画像path,
                        techpack: {寸法行たち, 素材, 原価, 納期}}, ...]
    @return {'ok': bool, 'パス': str, '枚数': int} または {'ok': False, '訳': str}
    """
    if not 商品リスト:
        return {'ok': False, '訳': '商品が1件もありません'}

    prs = Presentation()
    prs.slide_width = スライド幅
    prs.slide_height = スライド高さ

    _表紙を作る(prs, タイトル)
    for 商品 in 商品リスト:
        _商品スライドを作る(prs, 商品)

    os.makedirs(os.path.dirname(出力パス), exist_ok=True)
    prs.save(出力パス)
    return {'ok': True, 'パス': 出力パス, '枚数': len(商品リスト) + 1}
